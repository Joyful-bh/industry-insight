"""Refresh an exported track pool using existing evidence, without model calls or DB writes."""

import argparse
import csv
import json
import shutil
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook
from sqlalchemy.orm import Session

from scripts.event_stage_supervisor import save
from scripts.export_market_pilot import export_pilot
from track_insight.infrastructure.database import get_engine
from track_insight.tracks.contracts import CandidateTrackOutput
from track_insight.tracks.evidence_strength import SCORING_VERSION, recalculate_evidence_strength


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("directory", type=Path)
    args = parser.parse_args()
    root = args.directory
    path = root / "tracks.json"
    data = json.loads(path.read_text(encoding="utf-8"))
    tracks = [CandidateTrackOutput.model_validate(t) for t in data["tracks"]]
    engine = get_engine()
    try:
        with Session(engine) as session:
            audit = recalculate_evidence_strength(session, tracks)
    finally:
        engine.dispose()
    if any(row["missing_event_ids"] for row in audit):
        raise ValueError("Missing cited events: do not overwrite results against an incomplete DB")
    backup = root / "status_backup"
    backup.mkdir(exist_ok=True)
    for name in ("tracks.json", "status.json", "归并结果对照.xlsx", "赛道结果.csv"):
        if (root / name).exists() and not (backup / name).exists():
            shutil.copy2(root / name, backup / name)
    data["tracks"] = [t.model_dump(mode="json") for t in tracks]
    data["status_rule_version"] = SCORING_VERSION
    save(path, data)
    save(root / "evidence_status_audit.json", audit)
    summary = dict(
        version=SCORING_VERSION,
        before=dict(Counter(row["previous"]["status"] for row in audit)),
        after=dict(Counter(t.status for t in tracks)),
        transitions=dict(
            Counter(row["previous"]["status"] + "->" + row["status"] for row in audit)
        ),
        reasons=dict(Counter(row["reason"] for row in audit)),
        confidence_unchanged=True,
        semantic_coverage_rechecked=False,
    )
    save(root / "status_refresh.json", summary)
    status_path = root / "status.json"
    if status_path.exists():
        state = json.loads(status_path.read_text(encoding="utf-8"))
        state["evidence_status_refresh"] = summary
        save(status_path, state)
    export_pilot(root)
    csv_path = root / "赛道结果.csv"
    with csv_path.open(encoding="utf-8-sig", newline="") as stream:
        rows = list(csv.reader(stream))
    assert len(rows) == len(tracks) + 1
    headers = ["状态", "判定依据", "去重证据单元数"]
    rows[0].extend(headers)
    for row, track, decision in zip(rows[1:], tracks, audit, strict=True):
        assert row[3] == track.name and row[4] == track.definition
        row.extend([track.status, decision["reason"], str(decision["independent_evidence_units"])])
    with csv_path.open("w", encoding="utf-8-sig", newline="") as stream:
        csv.writer(stream).writerows(rows)
    book = load_workbook(root / "归并结果对照.xlsx")
    sheet = book["赛道结果"]
    for offset, title in enumerate(headers, 9):
        sheet.cell(1, offset, title)
        sheet.column_dimensions[sheet.cell(1, offset).column_letter].width = 30
    for index, row in enumerate(rows[1:], 2):
        for column, value in enumerate(row[8:], 9):
            sheet.cell(index, column, value)
    sheet.auto_filter.ref = sheet.dimensions
    temporary = root / ".status_view.xlsx"
    book.save(temporary)
    book.close()
    temporary.replace(root / "归并结果对照.xlsx")
    print(json.dumps(summary, ensure_ascii=False))


if __name__ == "__main__":
    main()
