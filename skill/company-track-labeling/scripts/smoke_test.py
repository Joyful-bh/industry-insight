"""Exercise the full API workflow against a deterministic local test service."""

import argparse
import csv
import hashlib
import json
import os
import subprocess
import sys
import threading
from collections import Counter
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path

from common import SKILL
from openpyxl import Workbook, load_workbook
from openpyxl.styles import PatternFill


class MockService(BaseHTTPRequestHandler):
    calls = Counter()
    failures = []

    def log_message(self, *_):
        pass

    def do_POST(self):
        try:
            body = json.loads(self.rfile.read(int(self.headers["Content-Length"])))
            if self.path.endswith("/embeddings"):
                self.calls["embedding"] += 1
                dimension = body["dimensions"]
                data = []
                for i, text in enumerate(body["input"]):
                    digest = hashlib.sha256(text.encode()).digest()
                    vector = [(digest[j % len(digest)] + 1) / 256 for j in range(dimension)]
                    data.append({"index": i, "embedding": vector})
                result = {"data": data, "usage": {"total_tokens": 10}}
            else:
                messages = body.get("messages") or body["input"]
                prompt = messages[-1]["content"]
                if "候选企业companies：" in prompt:
                    self.calls["local_verification"] += 1
                    companies = json.loads(prompt.split("候选企业companies：", 1)[1])
                    decisions = []
                    for company in companies:
                        cid = company["candidate_id"]
                        match = cid.split("::")[0].endswith("1")
                        exclude = cid.split("::")[0].endswith("4")
                        decisions.append(
                            {
                                "candidate_id": cid,
                                "label": "MATCH"
                                if match
                                else "NOT_MATCH"
                                if exclude
                                else "EVIDENCE_INSUFFICIENT",
                                "evidence_sufficient": match or exclude,
                                "confidence": "HIGH",
                                "reason_code": "CORE_PRODUCT_ACTIVITY"
                                if match
                                else "UNRELATED_BUSINESS"
                                if exclude
                                else "MISSING_COMPANY_INFORMATION",
                                "brief_reason": "固定测试响应，不用于实际企业判断。",
                            }
                        )
                    answer = {"decisions": decisions}
                elif "企业身份与已有资料：" in prompt:
                    advanced = self.path.endswith("/responses")
                    self.calls["advanced_search" if advanced else "basic_search"] += 1
                    company = json.loads(prompt.split("企业身份与已有资料：", 1)[1])
                    insufficient = company["company_id"].endswith("3") and not advanced
                    answer = {
                        "company_id": company["company_id"],
                        "information_status": "INSUFFICIENT" if insufficient else "SUFFICIENT",
                        "supplementary_profile": {
                            key: [] if insufficient else ["固定测试业务材料"]
                            for key in (
                                "core_business",
                                "products_services",
                                "actual_activities",
                                "manufacturing_rd_status",
                                "application_markets",
                            )
                        },
                    }
                else:
                    self.calls["adjudication"] += 1
                    data = json.loads(prompt.split("输入：", 1)[1])
                    answer = {
                        "track_id": data["track"]["track_id"],
                        "decisions": [
                            {
                                "candidate_id": c["candidate_id"],
                                "company_id": c["company_id"],
                                "label": "MATCH",
                                "brief_reason": "固定测试复判。",
                            }
                            for c in data["companies"]
                        ],
                    }
                if self.path.endswith("/responses"):
                    result = {
                        "output": [
                            {"type": "web_search_call"},
                            {
                                "type": "message",
                                "content": [{"type": "output_text", "text": json.dumps(answer)}],
                            },
                        ],
                        "usage": {"total_tokens": 10},
                    }
                else:
                    result = {
                        "choices": [{"message": {"content": json.dumps(answer)}}],
                        "usage": {"total_tokens": 10},
                    }
            payload = json.dumps(result).encode()
            self.send_response(200)
        except Exception as error:
            self.failures.append(str(error))
            payload = json.dumps({"error": str(error)}).encode()
            self.send_response(500)
        self.send_header("Content-Type", "application/json")
        self.send_header("Content-Length", str(len(payload)))
        self.end_headers()
        self.wfile.write(payload)


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--workspace", type=Path, required=True)
    parser.add_argument("--reference-project", type=Path)
    parser.add_argument("--dependencies", type=Path)
    args = parser.parse_args()
    workspace = args.workspace.resolve()
    if workspace.exists() and any(workspace.iterdir()):
        raise ValueError("Use an empty test workspace")
    workspace.mkdir(parents=True, exist_ok=True)
    caller = workspace / "无关启动目录"
    caller.mkdir()
    companies = workspace / "企业 样例.xlsx"
    tracks = workspace / "赛道 样例.jsonl"
    book = Workbook()
    sheet = book.active
    sheet.title = "企业"
    sheet.append(["企业名称", "统一社会信用代码", "经营状态", "经营范围", "企业简介"])
    sheet["A1"].fill = PatternFill("solid", fgColor="FF112233")
    for i in range(1, 5):
        sheet.append(
            [f"测试企业{i}", f"TEST{i}", "存续", "机器人关节研发生产", "机器人关节生产业务"]
        )
    book.create_sheet("说明")["A1"] = "原工作表须保留"
    book.save(companies)
    tracks.write_text(
        json.dumps(
            {
                "track_id": "TEST-TRACK",
                "name": "机器人关节制造",
                "definition": "研发与制造机器人关节模组",
                "core_products_services": ["机器人关节"],
                "included_activities": ["机器人关节研发生产"],
                "excluded_activities": ["仅销售"],
                "observable_company_features": ["关节生产线"],
            },
            ensure_ascii=False,
        )
        + "\n",
        encoding="utf-8",
    )
    input_hash = hashlib.sha256(companies.read_bytes()).hexdigest()
    server = ThreadingHTTPServer(("127.0.0.1", 0), MockService)
    thread = threading.Thread(target=server.serve_forever, daemon=True)
    thread.start()
    env = os.environ.copy()
    env["PYTHONIOENCODING"] = "utf-8"
    env["PYTHONDONTWRITEBYTECODE"] = "1"
    for key in list(env):
        if key.startswith("DASHSCOPE_"):
            env.pop(key)
    failed_output = ""

    def run(command):
        nonlocal failed_output
        result = subprocess.run(
            [sys.executable, "-B", *map(str, command)],
            cwd=caller,
            env=env,
            capture_output=True,
            text=True,
            encoding="utf-8",
        )
        if result.returncode:
            failed_output = result.stdout + result.stderr
            raise RuntimeError(failed_output)

    try:
        user = workspace / "用户工作目录"
        run([SKILL / "scripts/setup.py", "init", "--workspace", user])
        # This test can use already-installed test-host dependencies.
        if args.dependencies:
            import shutil

            shutil.copytree(args.dependencies.resolve(), user / ".python-deps", dirs_exist_ok=True)
        path = user / ".env"
        path.write_text(
            path.read_text(encoding="utf-8")
            .replace("DASHSCOPE_API_KEY=", "DASHSCOPE_API_KEY=test-only-local-key")
            .replace(
                "https://dashscope.aliyuncs.com/compatible-mode/v1",
                f"http://127.0.0.1:{server.server_port}/v1",
            ),
            encoding="utf-8",
        )
        wrapper = SKILL / "scripts/run_pipeline.py"
        run(
            [
                wrapper,
                "run",
                "--workspace",
                user,
                "--companies",
                companies,
                "--tracks",
                tracks,
                "--candidate-budget",
                "4",
                "--core-budget",
                "4",
                "--advanced-search",
            ]
        )
        state_path = user / ".run-state.json"
        state = json.loads(state_path.read_text(encoding="utf-8"))
        assert state["status"] == "completed"
        calls_before = dict(MockService.calls)
        state["status"] = "interrupted"
        state_path.write_text(json.dumps(state), encoding="utf-8")
        run([wrapper, "resume", "--workspace", user])
        assert dict(MockService.calls) == calls_before, "Resume should reuse model caches"
        output = load_workbook(state["options"]["output"])
        assert output.sheetnames == ["企业", "说明"]
        assert output["企业"]["A1"].fill.fgColor.rgb == "FF112233"
        assert output["企业"].max_row == 5
        assert "最终赛道标签" in [c.value for c in output["企业"][1]]
        assert hashlib.sha256(companies.read_bytes()).hexdigest() == input_hash
        assert not MockService.failures, MockService.failures
        for stage in (
            "embedding",
            "local_verification",
            "basic_search",
            "advanced_search",
            "adjudication",
        ):
            assert MockService.calls[stage] > 0, stage
        if args.reference_project:
            reference = workspace / "reference"
            run(
                [
                    args.reference_project.resolve() / "scripts/run_company_track_pipeline.py",
                    "--companies",
                    companies,
                    "--tracks",
                    tracks,
                    "--env",
                    path,
                    "--work-dir",
                    reference,
                    "--output",
                    workspace / "reference.xlsx",
                    "--candidate-budget",
                    "4",
                    "--core-budget",
                    "4",
                    "--advanced-search",
                ]
            )
            for relative in (
                "candidates/candidates_hybrid.csv",
                "web_basic/candidate_decisions_final.csv",
                "web_advanced/candidate_decisions_final.csv",
            ):

                def rows(p):
                    with p.open(encoding="utf-8-sig", newline="") as handle:
                        return list(csv.DictReader(handle))

                assert rows(user / "work" / relative) == rows(reference / relative), relative
            comparison = load_workbook(workspace / "reference.xlsx")
            assert list(comparison["企业"].values) == list(output["企业"].values)
        print(
            json.dumps(
                {
                    "offline_full_pipeline": True,
                    "advanced_search": True,
                    "cache_resume": True,
                    "workbook_preserved": True,
                    "reference_equivalent": bool(args.reference_project),
                    "api_calls": dict(MockService.calls),
                }
            )
        )
        return 0
    finally:
        server.shutdown()
        server.server_close()


if __name__ == "__main__":
    raise SystemExit(main())
