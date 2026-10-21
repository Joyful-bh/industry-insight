#!/usr/bin/env python3
"""Prepare normalized company records and classification text from a regional subset."""

from __future__ import annotations

import argparse
import csv
import json
import re
from collections import Counter
from pathlib import Path
from typing import Any

import pyarrow as pa
import pyarrow.parquet as pq
from openpyxl import load_workbook


ACTIVE_STATUSES = {"存续", "在业", "开业", "正常", "在营（开业）企业", "在营", "开业登记"}
INACTIVE_PREFIXES = ("注销", "吊销", "迁出", "撤销", "清算", "停业")
TEXT_FIELDS = ("企业名称", "主行业", "国标行业领域", "经营范围", "企业简介")
REGISTRY_PROFILE_MARKERS = (
    "成立于", "法定代表人为", "注册资本为", "统一社会信用代码为", "企业注册地址位于",
    "所属行业为", "经营范围包含", "企业当前经营状态为",
)


def clean_text(value: Any) -> str:
    text = "" if value is None else str(value)
    text = text.replace("\u3000", " ").replace("\xa0", " ")
    return re.sub(r"\s+", " ", text).strip()


def normalize_company_name(value: Any) -> str:
    text = clean_text(value).lower()
    return re.sub(r"[\s·•・,，.。()（）\[\]【】]", "", text)


def normalize_status(value: Any) -> str:
    status = clean_text(value)
    if not status:
        return "UNKNOWN"
    if status in ACTIVE_STATUSES:
        return "ACTIVE"
    if status.startswith(INACTIVE_PREFIXES):
        return "INACTIVE"
    return "UNKNOWN"


def is_registry_boilerplate_profile(value: Any) -> bool:
    text = clean_text(value)
    return sum(marker in text for marker in REGISTRY_PROFILE_MARKERS) >= 3


def make_classification_text(record: dict[str, str]) -> str:
    labels = {
        "企业名称": "企业名称", "主行业": "主行业", "国标行业领域": "国标行业",
        "经营范围": "经营范围", "企业简介": "企业简介",
    }
    parts = [f"{labels[field]}：{record[field]}" for field in TEXT_FIELDS if record.get(field)]
    return "\n".join(parts)


def prepare_record(raw: dict[str, Any]) -> dict[str, Any]:
    values = {key: clean_text(value) for key, value in raw.items()}
    status = normalize_status(values.get("经营状态"))
    flags = []
    if status == "UNKNOWN": flags.append("UNKNOWN_STATUS")
    if status == "INACTIVE": flags.append("INACTIVE_COMPANY")
    if not values.get("经营范围"): flags.append("MISSING_BUSINESS_SCOPE")
    if not values.get("企业简介"): flags.append("MISSING_COMPANY_PROFILE")
    profile_for_classification = values.get("企业简介", "")
    if is_registry_boilerplate_profile(profile_for_classification):
        flags.append("REGISTRY_BOILERPLATE_PROFILE")
        profile_for_classification = ""
    if not values.get("主行业") and not values.get("国标行业领域"): flags.append("MISSING_INDUSTRY")
    classification_values = dict(values)
    classification_values["企业简介"] = profile_for_classification
    text = make_classification_text(classification_values)
    if not text: flags.append("EMPTY_CLASSIFICATION_TEXT")
    return {
        "company_id": values.get("统一社会信用代码") or f"source_row:{values.get('source_row_number', '')}",
        "company_name": values.get("企业名称", ""),
        "company_name_normalized": normalize_company_name(values.get("企业名称")),
        "unified_social_credit_code": values.get("统一社会信用代码", ""),
        "province": values.get("省份", ""),
        "city": values.get("城市", ""),
        "district": values.get("区县", ""),
        "registration_status_raw": values.get("经营状态", ""),
        "registration_status": status,
        "classification_eligible": status != "INACTIVE" and bool(text),
        "primary_industry": values.get("主行业", ""),
        "national_industry_code": values.get("国标行业领域代码", ""),
        "national_industry_name": values.get("国标行业领域", ""),
        "national_industry_middle_code": values.get("国标行业中类代码", ""),
        "business_scope": values.get("经营范围", ""),
        "company_profile": values.get("企业简介", ""),
        "company_profile_for_classification": profile_for_classification,
        "former_name": values.get("曾用名", ""),
        "established_date": values.get("成立日期", ""),
        "employee_count_raw": values.get("员工人数", ""),
        "revenue_raw": values.get("营业收入", ""),
        "customer_level": values.get("客户分级", ""),
        "classification_text": text,
        "data_quality_flags": flags,
        "source_row_number": values.get("source_row_number", ""),
        "region_match_fields": values.get("region_match_fields", ""),
    }


def load_source(path: Path, sheet_name: str | None) -> list[dict[str, Any]]:
    suffix = path.suffix.lower()
    if suffix == ".parquet":
        return pq.read_table(path).to_pylist()
    if suffix == ".csv":
        with path.open("r", encoding="utf-8-sig", newline="") as handle:
            return list(csv.DictReader(handle))
    if suffix == ".xlsx":
        workbook = load_workbook(path, read_only=True, data_only=True)
        if sheet_name and sheet_name not in workbook.sheetnames:
            workbook.close()
            raise ValueError(f"worksheet not found: {sheet_name}")
        sheet = workbook[sheet_name] if sheet_name else workbook[workbook.sheetnames[0]]
        rows = sheet.iter_rows(values_only=True)
        try:
            headers = [clean_text(value) for value in next(rows)]
        except StopIteration as exc:
            workbook.close()
            raise ValueError("company workbook is empty") from exc
        if not any(headers):
            workbook.close()
            raise ValueError("company workbook has no header row")
        records = [dict(zip(headers, row)) for row in rows if any(value is not None for value in row)]
        workbook.close()
        return records
    raise ValueError(f"unsupported company input format: {suffix}")


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input", type=Path, required=True)
    parser.add_argument("--sheet", help="Worksheet name for XLSX input; defaults to the first sheet")
    parser.add_argument("--output-dir", type=Path, default=Path("data/processed/company_track_pipeline/companies"))
    args = parser.parse_args()
    args.output_dir.mkdir(parents=True, exist_ok=True)

    source = load_source(args.input, args.sheet)
    required = {"企业名称", "统一社会信用代码"}
    missing = required - set(source[0] if source else {})
    if missing:
        raise ValueError(f"company input is missing required columns: {sorted(missing)}")
    records = []
    for row_number, row in enumerate(source, start=2):
        normalized_source = dict(row)
        if not clean_text(normalized_source.get("source_row_number")):
            normalized_source["source_row_number"] = str(row_number)
        records.append(prepare_record(normalized_source))
    ids = [x["company_id"] for x in records]
    duplicate_ids = [key for key, count in Counter(ids).items() if count > 1]
    if duplicate_ids:
        raise ValueError(f"Duplicate company IDs: {duplicate_ids[:10]}")

    parquet_path = args.output_dir / "companies_classification_ready.parquet"
    jsonl_path = args.output_dir / "companies_classification_ready.jsonl"
    csv_path = args.output_dir / "companies_classification_ready.csv"
    table = pa.Table.from_pylist(records)
    pq.write_table(table, parquet_path, compression="zstd")
    with jsonl_path.open("w", encoding="utf-8", newline="\n") as handle:
        for record in records:
            handle.write(json.dumps(record, ensure_ascii=False, separators=(",", ":")) + "\n")
    csv_fields = [x for x in records[0] if x != "data_quality_flags"] + ["data_quality_flags"]
    with csv_path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=csv_fields)
        writer.writeheader()
        for record in records:
            row = dict(record)
            row["data_quality_flags"] = ";".join(record["data_quality_flags"])
            writer.writerow(row)

    statuses = Counter(x["registration_status"] for x in records)
    flags = Counter(flag for x in records for flag in x["data_quality_flags"])
    report = {
        "source_file": str(args.input), "record_count": len(records),
        "classification_eligible_count": sum(x["classification_eligible"] for x in records),
        "status_counts": dict(statuses), "quality_flag_counts": dict(flags),
        "duplicate_company_id_count": len(duplicate_ids),
        "classification_text_fields": list(TEXT_FIELDS),
        "outputs": {"parquet": str(parquet_path), "jsonl": str(jsonl_path), "csv": str(csv_path)},
    }
    (args.output_dir / "preparation_report.json").write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(report, ensure_ascii=False))


if __name__ == "__main__":
    main()
