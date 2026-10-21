#!/usr/bin/env python3
"""Merge final company-track decisions into a human-readable company workbook."""

from __future__ import annotations

import argparse
import csv
from collections import defaultdict
from copy import copy
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook


PROFILE_LABELS = {
    "SUFFICIENT": "充分",
    "PARTIAL": "部分",
    "INSUFFICIENT": "不足",
}


def read_csv(path: Path) -> list[dict[str, str]]:
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        return list(csv.DictReader(handle))


def read_company_table(
    path: Path, sheet_name: str | None = None
) -> tuple[list[str], list[tuple[Any, ...]]]:
    if path.suffix.lower() == ".csv":
        with path.open("r", encoding="utf-8-sig", newline="") as handle:
            reader = csv.reader(handle)
            headers = next(reader)
            return headers, [tuple(row) for row in reader]
    if path.suffix.lower() == ".xlsx":
        workbook = load_workbook(path, read_only=True, data_only=True)
        if sheet_name and sheet_name not in workbook.sheetnames:
            workbook.close()
            raise ValueError(f"worksheet not found: {sheet_name}")
        sheet = workbook[sheet_name] if sheet_name else workbook[workbook.sheetnames[0]]
        headers = [cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
        rows = list(sheet.iter_rows(min_row=2, values_only=True))
        workbook.close()
        return [str(header or "") for header in headers], rows
    raise ValueError(f"unsupported company table format: {path.suffix}")


def distinct(values: list[str]) -> list[str]:
    seen: set[str] = set()
    output: list[str] = []
    for value in values:
        value = str(value or "").strip()
        if value and value not in seen:
            seen.add(value)
            output.append(value)
    return output


def effective_profile_status(row: dict[str, str]) -> str:
    return (
        row.get("advanced_web_information_status")
        or row.get("web_information_status")
        or ""
    )


def effective_reason(row: dict[str, str]) -> str:
    return (
        row.get("advanced_web_recheck_brief_reason")
        or row.get("web_recheck_brief_reason")
        or row.get("verification_brief_reason")
        or ""
    ).strip()


def summarize_company(rows: list[dict[str, str]]) -> dict[str, Any]:
    ordered = sorted(rows, key=lambda row: int(row.get("global_rank") or 10**9))
    confirmed = distinct([row["track_name"] for row in ordered if row["decision_status"] == "CONFIRMED_MATCH"])
    probable = distinct([row["track_name"] for row in ordered if row["decision_status"] == "PROBABLE_MATCH"])
    unresolved = distinct([row["track_name"] for row in ordered if row["decision_status"] == "UNRESOLVED"])
    excluded = [row for row in ordered if row.get("automatic_disposition") == "EXCLUDE"]
    accepted = [row for row in ordered if row.get("automatic_disposition") == "ACCEPT"]

    status_values = distinct([effective_profile_status(row) for row in ordered])
    profile_status = "、".join(PROFILE_LABELS.get(value, value) for value in status_values)
    final_labels = [f"{name}（确认）" for name in confirmed] + [f"{name}（较可能）" for name in probable]

    conclusion_parts: list[str] = []
    if confirmed:
        conclusion_parts.append("确认匹配：" + "、".join(confirmed))
    if probable:
        conclusion_parts.append("较可能匹配：" + "、".join(probable))
    if unresolved:
        conclusion_parts.append("仍待确认：" + "、".join(unresolved))
    if not accepted and not unresolved:
        conclusion_parts.append("本轮候选均已排除，暂无可用赛道标签")

    reasons = distinct([
        f"{row['track_name']}：{effective_reason(row)}"
        for row in ordered
        if row.get("automatic_disposition") == "ACCEPT" or row["decision_status"] == "UNRESOLVED"
    ])
    return {
        "是否进入候选召回": "是",
        "最终赛道标签": "；".join(final_labels),
        "确认匹配赛道": "、".join(confirmed),
        "较可能匹配赛道": "、".join(probable),
        "待确认赛道": "、".join(unresolved),
        "企业补充档案状态": profile_status,
        "候选赛道数": len(ordered),
        "接受赛道数": len(accepted),
        "排除赛道数": len(excluded),
        "待确认赛道数": len(unresolved),
        "简单判断": "；".join(conclusion_parts) + "。",
        "判断依据摘要": "；".join(reasons),
    }


def export(args: argparse.Namespace) -> dict[str, int]:
    decisions = read_csv(args.decisions)
    by_id: dict[str, list[dict[str, str]]] = defaultdict(list)
    for row in decisions:
        by_id[str(row["company_id"]).strip()].append(row)

    source_headers, source_rows = read_company_table(args.input, args.sheet)
    source_indexes = {str(header): index for index, header in enumerate(source_headers)}
    for required_header in ("企业名称", "统一社会信用代码"):
        if required_header not in source_indexes:
            raise ValueError(f"source workbook is missing required column: {required_header}")

    result_headers = [
        "是否进入候选召回",
        "最终赛道标签",
        "确认匹配赛道",
        "较可能匹配赛道",
        "待确认赛道",
        "企业补充档案状态",
        "候选赛道数",
        "接受赛道数",
        "排除赛道数",
        "待确认赛道数",
        "简单判断",
        "判断依据摘要",
    ]

    if args.input.suffix.lower() == ".xlsx":
        workbook = load_workbook(args.input)
        company_sheet = workbook[args.sheet] if args.sheet else workbook[workbook.sheetnames[0]]
        workbook_headers = [
            str(company_sheet.cell(1, column).value or "")
            for column in range(1, company_sheet.max_column + 1)
        ]
        existing = [
            workbook_headers.index(header) + 1
            for header in result_headers
            if header in workbook_headers
        ]
        start_col = min(existing) if existing else company_sheet.max_column + 1
    else:
        workbook = Workbook()
        company_sheet = workbook.active
        company_sheet.title = "Sheet1"
        for column, header in enumerate(source_headers, start=1):
            company_sheet.cell(1, column, header)
        for row_index, source_row in enumerate(source_rows, start=2):
            for column, value in enumerate(source_row, start=1):
                company_sheet.cell(row_index, column, value)
        start_col = len(source_headers) + 1

    header_template = company_sheet.cell(1, max(start_col - 1, 1))
    for offset, header in enumerate(result_headers):
        cell = company_sheet.cell(1, start_col + offset, header)
        cell.font = copy(header_template.font)
        cell.fill = copy(header_template.fill)
        cell.border = copy(header_template.border)
        cell.alignment = copy(header_template.alignment)
        cell.number_format = header_template.number_format
        cell.protection = copy(header_template.protection)

    matched_company_ids: set[str] = set()
    for row_index, source_row in enumerate(source_rows, start=2):
        company_id = str(source_row[source_indexes["统一社会信用代码"]] or "").strip()
        if not company_id:
            source_number = (
                str(source_row[source_indexes["source_row_number"]] or "").strip()
                if "source_row_number" in source_indexes
                else str(row_index)
            )
            company_id = f"source_row:{source_number}"
        rows = by_id.get(company_id, [])
        if rows:
            summary = summarize_company(rows)
            matched_company_ids.add(company_id)
        else:
            summary = {header: "" for header in result_headers}
            summary.update({
                "是否进入候选召回": "否",
                "候选赛道数": 0,
                "接受赛道数": 0,
                "排除赛道数": 0,
                "待确认赛道数": 0,
                "简单判断": "本轮未进入候选召回，暂无赛道标签；不等同于确认不属于任何赛道。",
            })
        for offset, header in enumerate(result_headers):
            company_sheet.cell(row_index, start_col + offset, summary[header])

    final_column = start_col + len(result_headers) - 1
    if company_sheet.auto_filter.ref:
        company_sheet.auto_filter.ref = (
            f"A1:{company_sheet.cell(company_sheet.max_row, final_column).coordinate}"
        )
    for table in company_sheet.tables.values():
        table.ref = f"A1:{company_sheet.cell(company_sheet.max_row, final_column).coordinate}"

    args.output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(args.output)
    accepted_rows = [row for row in decisions if row.get("automatic_disposition") == "ACCEPT"]
    unresolved_rows = [row for row in decisions if row.get("automatic_disposition") == "UNRESOLVED"]
    return {
        "source_company_count": len(source_rows),
        "companies_with_candidates": len(matched_company_ids),
        "companies_without_candidates": len(source_rows) - len(matched_company_ids),
        "candidate_detail_count": len(decisions),
        "accepted_pair_count": len(accepted_rows),
        "accepted_company_count": len({row["company_id"] for row in accepted_rows}),
        "unresolved_pair_count": len(unresolved_rows),
    }


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input", type=Path, required=True)
    parser.add_argument("--sheet", help="Worksheet to update; defaults to the first sheet")
    parser.add_argument(
        "--decisions",
        type=Path,
        required=True,
    )
    parser.add_argument(
        "--output",
        type=Path,
        required=True,
    )
    return parser.parse_args()


if __name__ == "__main__":
    result = export(parse_args())
    print(result)
