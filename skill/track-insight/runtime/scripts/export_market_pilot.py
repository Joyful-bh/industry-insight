"""Export only the selected pilot results, with original candidates for comparison."""

import argparse
import csv
import json
from pathlib import Path


def export_pilot(directory):
    pilot = json.loads((directory / "pilot.json").read_text(encoding="utf-8"))
    preparation = json.loads((directory / "preparation.json").read_text(encoding="utf-8"))
    sources = {f"C{row['index']}": row for row in pilot["selected_tracks"]}
    packages = {
        ref: [
            index for index in pilot["package_ids"]
            if int(ref[1:]) in preparation["packages"][index]
        ]
        for ref in sources
    }
    headers = [
        "结果编号", "处理方式", "包号", "赛道名称", "范围定义",
        "原候选数", "原候选编号", "原候选名称"
    ]
    results, details, used = [], [], set()
    for index, group in enumerate(pilot["stats"]["merge_groups"], 1):
        refs = group["member_refs"]
        used.update(refs)
        result_id = f"M{index}"
        results.append([
            result_id, "合并", ", ".join(map(str, sorted({p for r in refs for p in packages[r]}))),
            group["name"], group["boundary"], len(refs), ", ".join(refs),
            "\n".join(sources[r]["name"] for r in refs),
        ])
        for ref in refs:
            details.append([
                result_id, group["name"], ref, sources[ref]["name"], sources[ref]["definition"]
            ])
    for ref, source in sources.items():
        if ref in used:
            continue
        results.append([
            ref, "原样保留", ", ".join(map(str, packages[ref])), source["name"],
            source["definition"], 1, ref, source["name"],
        ])
        details.append([ref, source["name"], ref, source["name"], source["definition"]])
    tables = [
        ("赛道结果", headers, results),
        ("成员对照", ["结果编号", "结果名称", "原候选编号", "原候选名称", "原候选定义"], details),
    ]
    for name, columns, rows in tables:
        with (directory / f"{name}.csv").open("w", encoding="utf-8-sig", newline="") as stream:
            writer = csv.writer(stream)
            writer.writerow(columns)
            writer.writerows(rows)
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font
    except ImportError:
        return directory / "赛道结果.csv"
    book = Workbook()
    book.remove(book.active)
    for name, columns, rows in tables:
        sheet = book.create_sheet(name)
        sheet.append(columns)
        for row in rows:
            sheet.append(row)
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(bold=True)
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
            sheet.row_dimensions[row[0].row].height = (
                min(400, max(100, len(str(row[-1].value).splitlines()) * 16 + 10))
                if name == "赛道结果" else 100
            )
        widths = [14, 16, 16, 42, 100, 12, 30, 55] if name == "赛道结果" else [14, 42, 16, 42, 100]
        for column, width in zip(sheet.columns, widths, strict=True):
            sheet.column_dimensions[column[0].column_letter].width = width
    path = directory / "归并结果对照.xlsx"
    book.save(path)
    return path


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("directory", type=Path)
    args = parser.parse_args()
    print(export_pilot(args.directory))


if __name__ == "__main__":
    main()
